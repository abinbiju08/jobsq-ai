from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from groq import Groq
from typing import List, Optional
import os, json, io
from dotenv import load_dotenv

load_dotenv()
router = APIRouter()
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))

class Experience(BaseModel):
    company: str
    role: str
    start: str
    end: str
    bullets: List[str]

class Education(BaseModel):
    institution: str
    degree: str
    field: str
    year: str
    gpa: Optional[str] = ""

class Project(BaseModel):
    name: str
    description: str
    tech: str

class ResumeData(BaseModel):
    template: str
    full_name: str
    email: str
    phone: str
    location: str
    linkedin: Optional[str] = ""
    github: Optional[str] = ""
    portfolio: Optional[str] = ""
    target_role: str
    years_experience: str
    summary: Optional[str] = ""
    experiences: List[Experience] = []
    education: List[Education] = []
    technical_skills: List[str] = []
    soft_skills: List[str] = []
    projects: List[Project] = []
    certifications: List[str] = []
    languages: List[str] = []

def clean_text(text: str) -> str:
    """Remove characters that reportlab/docx can't render — replaces bullets and dashes with ASCII equivalents"""
    if not text:
        return ""
    replacements = {
        '\u2022': '-',  # bullet •
        '\u2023': '-',  # triangular bullet
        '\u25aa': '-',  # black small square
        '\u25a0': '-',  # black square
        '\u25cf': '-',  # black circle
        '\u2013': '-',  # en dash
        '\u2014': '-',  # em dash
        '\u2018': "'",  # left single quote
        '\u2019': "'",  # right single quote
        '\u201c': '"',  # left double quote
        '\u201d': '"',  # right double quote
        '\u00a0': ' ',  # non-breaking space
        '\u2026': '...',# ellipsis
        '\ufffd': '',   # replacement character
        '\u200b': '',   # zero-width space
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    # Remove any remaining non-latin characters that can't render
    result = ''
    for ch in text:
        if ord(ch) < 128 or ord(ch) in range(160, 256):
            result += ch
        else:
            result += ' '
    return result.strip()

TEMPLATE_COLORS = {
    "modern":       (0,   196, 132),
    "professional": (30,  58,  95),
    "executive":    (26,  26,  46),
    "minimal":      (45,  55,  72),
    "creative":     (124, 58,  237),
    "healthcare":   (8,   145, 178),
}

def generate_pdf(resume: dict, template: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm, mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY

    r, g, b = TEMPLATE_COLORS.get(template, (0, 196, 132))
    accent = colors.Color(r/255, g/255, b/255)
    dark   = colors.Color(0.08, 0.08, 0.08)
    gray   = colors.Color(0.35, 0.35, 0.35)
    lgray  = colors.Color(0.55, 0.55, 0.55)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.8*cm, rightMargin=1.8*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm
    )

    W = A4[0] - 3.6*cm  # usable width

    story = []

    # ── HEADER ──────────────────────────────────────────────────
    # Name
    story.append(Paragraph(
        f'<font size="22" color="#{r:02X}{g:02X}{b:02X}"><b>{clean_text(resume.get("full_name",""))}</b></font>',
        ParagraphStyle('name', fontName='Helvetica-Bold', fontSize=22, leading=26, spaceAfter=2)
    ))
    # Role
    story.append(Paragraph(
        f'<font size="12" color="#505050">{resume.get("target_role","")}</font>',
        ParagraphStyle('role', fontName='Helvetica', fontSize=12, leading=15, spaceAfter=4)
    ))
    # Contact line
    contact_parts = [resume.get("email",""), resume.get("phone",""), resume.get("location",""),
                     resume.get("linkedin",""), resume.get("github","")]
    contact_str = "  |  ".join([clean_text(c) for c in contact_parts if c])
    story.append(Paragraph(
        f'<font size="9" color="#666666">{contact_str}</font>',
        ParagraphStyle('contact', fontName='Helvetica', fontSize=9, leading=12, spaceAfter=8)
    ))
    story.append(HRFlowable(width=W, thickness=2, color=accent, spaceAfter=10))

    def section_header(title):
        story.append(Spacer(1, 6))
        story.append(Paragraph(
            f'<font size="10.5" color="#{r:02X}{g:02X}{b:02X}"><b>{title.upper()}</b></font>',
            ParagraphStyle('sec', fontName='Helvetica-Bold', fontSize=10.5, leading=14, spaceAfter=2)
        ))
        story.append(HRFlowable(width=W, thickness=1, color=accent, spaceAfter=6))

    body_style = ParagraphStyle('body', fontName='Helvetica', fontSize=10, leading=14,
                                textColor=colors.Color(0.23,0.23,0.23), spaceAfter=2)
    bullet_style = ParagraphStyle('bullet', fontName='Helvetica', fontSize=10, leading=14,
                                  textColor=colors.Color(0.23,0.23,0.23), leftIndent=12,
                                  firstLineIndent=0, spaceAfter=2, bulletIndent=2)

    # ── SUMMARY ──────────────────────────────────────────────────
    summary = clean_text(resume.get("professional_summary", ""))
    if summary:
        section_header("Professional Summary")
        story.append(Paragraph(summary, ParagraphStyle('sum', fontName='Helvetica', fontSize=10,
                    leading=15, textColor=colors.Color(0.23,0.23,0.23), alignment=TA_JUSTIFY, spaceAfter=4)))

    # ── EXPERIENCE ────────────────────────────────────────────────
    experiences = resume.get("experience", [])
    if experiences:
        section_header("Work Experience")
        for exp in experiences:
            # Role + Date on same row using table
            role_para = Paragraph(f'<b>{clean_text(exp.get("role",""))}</b>',
                ParagraphStyle('er', fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=dark))
            date_str = f'{exp.get("start","")} – {exp.get("end","")}'
            date_para = Paragraph(f'<font size="9" color="#888888">{date_str}</font>',
                ParagraphStyle('ed', fontName='Helvetica', fontSize=9, leading=14,
                               textColor=lgray, alignment=TA_RIGHT))
            t = Table([[role_para, date_para]], colWidths=[W*0.68, W*0.32])
            t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                                   ('LEFTPADDING',(0,0),(-1,-1),0),
                                   ('RIGHTPADDING',(0,0),(-1,-1),0),
                                   ('TOPPADDING',(0,0),(-1,-1),2),
                                   ('BOTTOMPADDING',(0,0),(-1,-1),2)]))
            story.append(t)
            # Company in accent color
            story.append(Paragraph(
                f'<font size="10" color="#{r:02X}{g:02X}{b:02X}"><i>{clean_text(exp.get("company",""))}</i></font>',
                ParagraphStyle('ec', fontName='Helvetica-Oblique', fontSize=10, leading=13,
                               textColor=accent, spaceAfter=3)))
            # Bullets
            for bl in exp.get("bullets", []):
                if bl and bl.strip():
                    story.append(Paragraph(f'- {clean_text(bl)}', bullet_style))
            story.append(Spacer(1, 6))

    # ── EDUCATION ─────────────────────────────────────────────────
    educations = resume.get("education", [])
    if educations:
        section_header("Education")
        for edu in educations:
            deg_para = Paragraph(f'<b>{clean_text(edu.get("degree",""))} in {clean_text(edu.get("field",""))}</b>',
                ParagraphStyle('degr', fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=dark))
            year_para = Paragraph(f'<font size="9" color="#888888">{edu.get("year","")}</font>',
                ParagraphStyle('eyp', fontName='Helvetica', fontSize=9, leading=14,
                               textColor=lgray, alignment=TA_RIGHT))
            t = Table([[deg_para, year_para]], colWidths=[W*0.75, W*0.25])
            t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                                   ('LEFTPADDING',(0,0),(-1,-1),0),
                                   ('RIGHTPADDING',(0,0),(-1,-1),0),
                                   ('TOPPADDING',(0,0),(-1,-1),2),
                                   ('BOTTOMPADDING',(0,0),(-1,-1),2)]))
            story.append(t)
            gpa_text = f"  |  GPA: {edu.get('gpa','')}" if edu.get('gpa') else ""
            story.append(Paragraph(
                f'<font size="10" color="#666666">{clean_text(edu.get("institution",""))}{clean_text(gpa_text)}</font>',
                ParagraphStyle('ei', fontName='Helvetica', fontSize=10, leading=13,
                               textColor=gray, spaceAfter=6)))

    # ── SKILLS ────────────────────────────────────────────────────
    tech  = resume.get("technical_skills", [])
    soft  = resume.get("soft_skills", [])
    if tech or soft:
        section_header("Technical Skills")
        if tech:
            story.append(Paragraph("  |  ".join([clean_text(s) for s in tech]), body_style))
        if soft:
            story.append(Paragraph(
                f'<b>Soft Skills:</b>  {", ".join([clean_text(s) for s in soft])}',
                ParagraphStyle('ss', fontName='Helvetica', fontSize=10, leading=14,
                               textColor=colors.Color(0.23,0.23,0.23), spaceAfter=4)))

    # ── PROJECTS ──────────────────────────────────────────────────
    projects = [p for p in resume.get("projects", []) if p.get("name","").strip()]
    if projects:
        section_header("Projects")
        for proj in projects:
            story.append(Paragraph(f'<b>{clean_text(proj.get("name",""))}</b>',
                ParagraphStyle('pn', fontName='Helvetica-Bold', fontSize=11, leading=14,
                               textColor=dark, spaceAfter=2)))
            story.append(Paragraph(clean_text(proj.get("description","")), body_style))
            story.append(Paragraph(
                f'<font color="#{r:02X}{g:02X}{b:02X}"><b>Tech:</b></font>  {clean_text(proj.get("tech",""))}',
                ParagraphStyle('pt', fontName='Helvetica', fontSize=9.5, leading=13,
                               textColor=gray, spaceAfter=6)))

    # ── CERTIFICATIONS ────────────────────────────────────────────
    certs = resume.get("certifications", [])
    if certs:
        section_header("Certifications")
        for c in certs:
            story.append(Paragraph(f'- {clean_text(c)}', bullet_style))
        story.append(Spacer(1, 4))

    # ── LANGUAGES ─────────────────────────────────────────────────
    langs = resume.get("languages", [])
    if langs:
        section_header("Languages")
        story.append(Paragraph("  |  ".join([clean_text(s) for s in langs]), body_style))

    doc.build(story)
    buf.seek(0)
    return buf.read()


def generate_docx(resume: dict, template: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    r, g, b = TEMPLATE_COLORS.get(template, (0, 196, 132))

    def add_hr(doc, color_hex="CCCCCC"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), f'{r:02X}{g:02X}{b:02X}')
        pBdr.append(bottom)
        pPr.append(pBdr)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)

    # Name
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(clean_text(resume.get("full_name","")).upper())
    run.bold = True; run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(r, g, b)

    # Role
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(resume.get("target_role",""))
    run.font.size = Pt(12); run.font.color.rgb = RGBColor(80,80,80)

    # Contact
    contact_parts = [resume.get("email",""), resume.get("phone",""), resume.get("location",""),
                     resume.get("linkedin",""), resume.get("github","")]
    contact_str = "  |  ".join([clean_text(c) for c in contact_parts if c])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(contact_str)
    run.font.size = Pt(9); run.font.color.rgb = RGBColor(100,100,100)
    add_hr(doc)

    def section_heading(title):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title.upper())
        run.bold = True; run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(r, g, b)
        add_hr(doc)

    # Summary
    if resume.get("professional_summary"):
        section_heading("Professional Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(resume["professional_summary"])
        run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)

    # Experience
    if resume.get("experience"):
        section_heading("Work Experience")
        for exp in resume["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(clean_text(exp.get("role","")))
            r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = RGBColor(20,20,20)
            date = f"  {exp.get('start','')} – {exp.get('end','')}"
            r2 = p.add_run(date)
            r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(130,130,130)

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_text(exp.get("company","")))
            run.italic = True; run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(r, g, b)

            for bl in exp.get("bullets", []):
                if bl and bl.strip():
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.left_indent = Inches(0.2)
                    run = p.add_run(clean_text(bl))
                    run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)

    # Education
    if resume.get("education"):
        section_heading("Education")
        for edu in resume["education"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(f"{edu.get('degree','')} in {edu.get('field','')}")
            r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = RGBColor(20,20,20)
            r2 = p.add_run(f"  |  {edu.get('year','')}")
            r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(130,130,130)

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            gpa = f"  |  GPA: {edu.get('gpa','')}" if edu.get('gpa') else ""
            run = p.add_run(f"{edu.get('institution','')}{gpa}")
            run.font.size = Pt(10); run.font.color.rgb = RGBColor(100,100,100)

    # Skills
    tech = resume.get("technical_skills", [])
    soft = resume.get("soft_skills", [])
    if tech or soft:
        section_heading("Technical Skills")
        if tech:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("  |  ".join([clean_text(s) for s in tech]))
            run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)
        if soft:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run("Soft Skills:  ")
            r1.bold = True; r1.font.size = Pt(10); r1.font.color.rgb = RGBColor(20,20,20)
            r2 = p.add_run(", ".join([clean_text(s) for s in soft]))
            r2.font.size = Pt(10); r2.font.color.rgb = RGBColor(60,60,60)

    # Projects
    projects = [p for p in resume.get("projects", []) if p.get("name","").strip()]
    if projects:
        section_heading("Projects")
        for proj in projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
            run = p.add_run(proj.get("name",""))
            run.bold = True; run.font.size = Pt(11); run.font.color.rgb = RGBColor(20,20,20)

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(proj.get("description",""))
            run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run("Tech: "); r1.bold = True
            r1.font.size = Pt(9.5); r1.font.color.rgb = RGBColor(r,g,b)
            r2 = p.add_run(proj.get("tech",""))
            r2.font.size = Pt(9.5); r2.font.color.rgb = RGBColor(100,100,100)

    # Certifications
    if resume.get("certifications"):
        section_heading("Certifications")
        for c in resume["certifications"]:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(clean_text(c))
            run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)

    # Languages
    if resume.get("languages"):
        section_heading("Languages")
        p = doc.add_paragraph()
        run = p.add_run("  |  ".join([clean_text(s) for s in resume["languages"]]))
        run.font.size = Pt(10); run.font.color.rgb = RGBColor(60,60,60)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@router.post("/generate")
async def generate_resume(data: ResumeData):
    try:
        exp_text = ""
        for exp in data.experiences:
            bullets = "\n".join([f"- {b}" for b in exp.bullets if b.strip()])
            exp_text += f"\n{exp.role} at {exp.company} ({exp.start} - {exp.end})\n{bullets}\n"
        edu_text = ""
        for edu in data.education:
            edu_text += f"\n{edu.degree} in {edu.field} - {edu.institution} ({edu.year})"
            if edu.gpa: edu_text += f" | GPA: {edu.gpa}"
        proj_text = ""
        for proj in data.projects:
            proj_text += f"\n{proj.name}: {proj.description} | Tech: {proj.tech}"

        prompt = f"""You are an expert resume writer. Create an ATS-optimized resume.

Name: {data.full_name} | Role: {data.target_role} | Experience: {data.years_experience} years
Contact: {data.email} | {data.phone} | {data.location}
LinkedIn: {data.linkedin} | GitHub: {data.github}
EXPERIENCE:{exp_text}
EDUCATION:{edu_text}
TECHNICAL SKILLS: {', '.join(data.technical_skills)}
SOFT SKILLS: {', '.join(data.soft_skills)}
PROJECTS:{proj_text}
CERTIFICATIONS: {', '.join(data.certifications)}
LANGUAGES: {', '.join(data.languages)}
USER SUMMARY: {data.summary}

Return ONLY valid JSON no markdown:
{{"professional_summary":"2-3 powerful ATS-optimized sentences","experience":[{{"company":"","role":"","start":"","end":"","bullets":["quantified achievement","achievement","achievement"]}}],"education":[{{"institution":"","degree":"","field":"","year":"","gpa":""}}],"technical_skills":[],"soft_skills":[],"projects":[{{"name":"","description":"improved description","tech":""}}],"certifications":[],"languages":[],"ats_keywords":[]}}"""

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3, max_tokens=2000
        )
        raw = response.choices[0].message.content.strip()
        if '<think>' in raw: raw = raw.split('</think>')[-1].strip()
        if '```' in raw: raw = raw.split('```')[1].replace('json','').strip()
        start = raw.find('{'); end = raw.rfind('}')
        if start != -1 and end != -1: raw = raw[start:end+1]
        resume_content = json.loads(raw)
        full_resume = {**resume_content, "full_name":data.full_name, "email":data.email,
                      "phone":data.phone, "location":data.location, "linkedin":data.linkedin,
                      "github":data.github, "portfolio":data.portfolio,
                      "target_role":data.target_role, "template":data.template}
        return {"success": True, "resume": full_resume}
    except Exception as e:
        import traceback; traceback.print_exc()
        return {"success": False, "error": str(e)}


@router.post("/download-pdf")
async def download_pdf(data: dict):
    try:
        resume = data.get("resume", {})
        template = data.get("template", "modern")
        pdf_bytes = generate_pdf(resume, template)
        name = resume.get("full_name", "Resume").replace(" ", "_")
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{name}_Resume.pdf"'}
        )
    except Exception as e:
        import traceback; traceback.print_exc()
        return {"error": str(e)}


@router.post("/download-docx")
async def download_docx(data: dict):
    try:
        resume = data.get("resume", {})
        template = data.get("template", "modern")
        docx_bytes = generate_docx(resume, template)
        name = resume.get("full_name", "Resume").replace(" ", "_")
        return StreamingResponse(
            io.BytesIO(docx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{name}_Resume.docx"'}
        )
    except Exception as e:
        import traceback; traceback.print_exc()
        return {"error": str(e)}

# ── UPLOAD & CONVERT EXISTING RESUME ─────────────────────────────
from fastapi import UploadFile, File
import PyPDF2

@router.post("/extract-resume")
async def extract_resume(file: UploadFile = File(...)):
    """Extract info from uploaded resume PDF and return structured data"""
    try:
        content = await file.read()
        # Extract text from PDF
        text = ""
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            for page in reader.pages:
                text += page.extract_text() or ""
        except:
            text = content.decode("utf-8", errors="ignore")

        if not text.strip():
            return {"success": False, "error": "Could not read the resume file"}

        prompt = f"""Extract all information from this resume and return ONLY valid JSON:
{{
  "full_name": "extracted name",
  "email": "email@example.com",
  "phone": "+91 XXXXXXXXXX",
  "location": "City, Country",
  "linkedin": "linkedin url or empty",
  "github": "github url or empty",
  "portfolio": "portfolio url or empty",
  "target_role": "most recent job title or target role",
  "years_experience": "number as string",
  "summary": "professional summary if present",
  "experiences": [
    {{
      "company": "company name",
      "role": "job title",
      "start": "Mon YYYY",
      "end": "Mon YYYY or Present",
      "bullets": ["responsibility or achievement 1", "achievement 2", "achievement 3"]
    }}
  ],
  "education": [
    {{
      "institution": "university name",
      "degree": "degree type",
      "field": "field of study",
      "year": "graduation year",
      "gpa": "gpa if mentioned or empty"
    }}
  ],
  "technical_skills": ["skill1", "skill2"],
  "soft_skills": ["skill1", "skill2"],
  "projects": [
    {{
      "name": "project name",
      "description": "project description",
      "tech": "technologies used"
    }}
  ],
  "certifications": ["cert1", "cert2"],
  "languages": ["language1"]
}}

Resume text:
{text[:4000]}"""

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
            max_tokens=2000
        )

        raw = response.choices[0].message.content.strip()
        if '<think>' in raw:
            raw = raw.split('</think>')[-1].strip()
        if '```' in raw:
            raw = raw.split('```')[1].replace('json','').strip()
        start = raw.find('{'); end = raw.rfind('}')
        if start != -1 and end != -1:
            raw = raw[start:end+1]

        extracted = json.loads(raw)
        return {"success": True, "data": extracted}

    except Exception as e:
        import traceback; traceback.print_exc()
        return {"success": False, "error": str(e)}


# ── COVER LETTER GENERATION ──────────────────────────────────────
class CoverLetterRequest(BaseModel):
    resume: dict
    company: str
    job_title: str
    job_description: Optional[str] = ""
    highlight: Optional[str] = ""
    tone: str = "Professional"

class CoverLetterFromScratch(BaseModel):
    full_name: str
    role: str
    skills: str
    company: str
    job_title: str
    job_description: Optional[str] = ""
    highlight: Optional[str] = ""
    tone: str = "Professional"

TONE_INSTRUCTIONS = {
    "Professional": "formal, confident, and polished. Use strong action verbs and maintain a businesslike tone throughout.",
    "Enthusiastic": "warm, energetic, and passionate. Show genuine excitement for the role while remaining professional.",
    "Creative": "engaging and memorable. Use vivid language and a unique opening that stands out from generic letters.",
    "Concise": "brief and to the point. Maximum 3 short paragraphs. Every sentence must add value.",
}

def build_cover_letter_prompt(name, role, skills_text, company, job_title, job_description, highlight, tone, summary=""):
    tone_instruction = TONE_INSTRUCTIONS.get(tone, TONE_INSTRUCTIONS["Professional"])
    return f"""Write a {tone} cover letter for a job application. The tone should be {tone_instruction}

APPLICANT DETAILS:
Name: {name}
Current/Target Role: {role}
Key Skills: {skills_text}
Professional Summary: {summary}

TARGET JOB:
Company: {company}
Position: {job_title}
Job Description: {job_description or 'Not provided'}
Specific highlights to mention: {highlight or 'None'}

Write a complete, ready-to-send cover letter. Include:
- Professional greeting
- Strong opening paragraph (why this company and role)
- 1-2 middle paragraphs highlighting relevant experience and skills matching the job
- Closing paragraph with call to action
- Professional sign-off

Return ONLY the cover letter text, no JSON, no markdown, no extra commentary. Make it natural and human-sounding."""


@router.post("/generate-cover-letter")
async def generate_cover_letter(data: CoverLetterRequest):
    """Generate cover letter from an existing generated resume"""
    try:
        resume = data.resume
        name = resume.get("full_name", "")
        role = clean_text(resume.get("target_role", ""))
        skills = resume.get("technical_skills", []) + resume.get("soft_skills", [])
        skills_text = ", ".join(skills[:12])
        summary = resume.get("professional_summary", "")

        prompt = build_cover_letter_prompt(
            name, role, skills_text, data.company,
            data.job_title, data.job_description,
            data.highlight, data.tone, summary
        )

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
            max_tokens=1000
        )

        text = response.choices[0].message.content.strip()
        if '<think>' in text:
            text = text.split('</think>')[-1].strip()

        return {"success": True, "cover_letter": text, "name": name, "company": data.company, "job_title": data.job_title}

    except Exception as e:
        import traceback; traceback.print_exc()
        return {"success": False, "error": str(e)}


@router.post("/generate-cover-letter-scratch")
async def generate_cover_letter_scratch(data: CoverLetterFromScratch):
    """Generate cover letter from manual input or uploaded resume extraction"""
    try:
        prompt = build_cover_letter_prompt(
            data.full_name, data.role, data.skills,
            data.company, data.job_title,
            data.job_description, data.highlight, data.tone
        )

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
            max_tokens=1000
        )

        text = response.choices[0].message.content.strip()
        if '<think>' in text:
            text = text.split('</think>')[-1].strip()

        return {"success": True, "cover_letter": text}

    except Exception as e:
        import traceback; traceback.print_exc()
        return {"success": False, "error": str(e)}


@router.post("/download-cover-letter-pdf")
async def download_cover_letter_pdf(data: dict):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY

        text = data.get("cover_letter", "")
        name = data.get("name", "Applicant")
        company = data.get("company", "")
        job_title = data.get("job_title", "")

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
            leftMargin=2*cm, rightMargin=2*cm,
            topMargin=2*cm, bottomMargin=2*cm)

        story = []
        dark = colors.Color(0.1, 0.1, 0.1)
        gray = colors.Color(0.4, 0.4, 0.4)

        # Name header
        story.append(Paragraph(clean_text(name),
            ParagraphStyle('name', fontName='Helvetica-Bold', fontSize=18,
                           leading=22, textColor=dark, spaceAfter=4)))
        if company and job_title:
            story.append(Paragraph(clean_text(f"Application for {job_title} - {company}"),
                ParagraphStyle('sub', fontName='Helvetica', fontSize=11,
                               leading=14, textColor=gray, spaceAfter=8)))
        story.append(HRFlowable(width='100%', thickness=1,
                                color=colors.Color(0,0.77,0.51), spaceAfter=16))

        body_style = ParagraphStyle('body', fontName='Helvetica', fontSize=10.5,
                                    leading=16, textColor=dark, alignment=TA_JUSTIFY,
                                    spaceAfter=10)

        # Split paragraphs and render
        for para in clean_text(text).split('\n\n'):
            para = para.strip()
            if not para:
                continue
            for line in para.split('\n'):
                line = line.strip()
                if line:
                    story.append(Paragraph(clean_text(line), body_style))
            story.append(Spacer(1, 4))

        doc.build(story)
        buf.seek(0)
        safe_name = name.replace(" ", "_")
        return StreamingResponse(
            io.BytesIO(buf.read()),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{safe_name}_Cover_Letter.pdf"'}
        )
    except Exception as e:
        import traceback; traceback.print_exc()
        return {"error": str(e)}


@router.post("/download-cover-letter-docx")
async def download_cover_letter_docx(data: dict):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        text = data.get("cover_letter", "")
        name = data.get("name", "Applicant")
        company = data.get("company", "")
        job_title = data.get("job_title", "")

        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(2); section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)

        # Name
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(name)
        run.bold = True; run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(10, 10, 10)

        if company and job_title:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(f"Application for {job_title} — {company}")
            run.font.size = Pt(11); run.font.color.rgb = RGBColor(100, 100, 100)

        # HR line
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '00C483')
        pBdr.append(bottom)
        pPr.append(pBdr)

        # Body paragraphs
        for para in text.split('\n\n'):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(clean_text(para.replace('\n', ' ')))
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(20, 20, 20)

        buf = io.BytesIO()
        doc.save(buf); buf.seek(0)
        safe_name = name.replace(" ", "_")
        return StreamingResponse(
            io.BytesIO(buf.read()),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{safe_name}_Cover_Letter.docx"'}
        )
    except Exception as e:
        import traceback; traceback.print_exc()
        return {"error": str(e)}