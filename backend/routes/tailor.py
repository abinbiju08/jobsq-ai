from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from supabase import create_client
from groq import Groq
from pydantic import BaseModel
import PyPDF2, io, os, json, re, unicodedata
from dotenv import load_dotenv
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.platypus import KeepTogether

load_dotenv()
router = APIRouter()

supabase = create_client(os.getenv("SUPABASE_URL"), os.getenv("SUPABASE_SERVICE_KEY"))
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))


# ─────────────────────────────────────────────
# PDF TEXT + URL EXTRACTION
# ─────────────────────────────────────────────

def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from PDF using PyPDF2 + grab hyperlink URLs from annotations."""
    text = ""
    urls_found = []

    # 1. Extract text with PyPDF2
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"

            # Extract URLs from annotations (hyperlinks embedded in PDF)
            if "/Annots" in page:
                for annot in page["/Annots"]:
                    obj = annot.get_object()
                    if obj.get("/Subtype") == "/Link":
                        uri = obj.get("/A", {}).get("/URI", "")
                        if uri and uri not in urls_found:
                            urls_found.append(uri)
    except Exception as e:
        print(f"PyPDF2 error: {e}")

    # 2. Also try pymupdf for better URL extraction from hyperlinks
    try:
        import pymupdf
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            for link in page.get_links():
                uri = link.get("uri", "")
                if uri and uri not in urls_found:
                    urls_found.append(uri)
        doc.close()
    except Exception as e:
        print(f"pymupdf link extract error: {e}")

    # 3. Append found URLs to the text so AI can pick them up
    if urls_found:
        text += "\n\nEXTRACTED_LINKS:\n"
        for url in urls_found:
            text += url + "\n"
            # Label them for the AI
            if "linkedin" in url.lower():
                text += f"LinkedIn: {url}\n"
            elif "github" in url.lower():
                text += f"GitHub: {url}\n"

    # 4. Also scan text for URLs using regex (catches plain-text URLs)
    url_pattern = r'(https?://[^\s\)\]>\"]+|linkedin\.com/in/[^\s\)\]>\"]+|github\.com/[^\s\)\]>\"]+)'
    text_urls = re.findall(url_pattern, text)
    for u in text_urls:
        if u not in urls_found:
            urls_found.append(u)
            if "linkedin" in u.lower():
                text += f"\nLinkedIn: {u}"
            elif "github" in u.lower():
                text += f"\nGitHub: {u}"

    print(f"Resume text length: {len(text)}")
    print(f"Resume text preview: {text[:600]}")
    print(f"URLs found: {urls_found}")
    return text.strip()


# ─────────────────────────────────────────────
# SKILL FILTERING
# ─────────────────────────────────────────────

TECH_WHITELIST = {
    'python','java','javascript','typescript','react','node','nodejs','angular','vue',
    'html','css','sql','mysql','postgresql','mongodb','redis','docker','kubernetes',
    'aws','gcp','azure','git','github','gitlab','linux','bash','shell','fastapi',
    'django','flask','express','spring','springboot','hibernate','rest','graphql',
    'tensorflow','pytorch','scikit','pandas','numpy','matplotlib','selenium','jest',
    'junit','postman','swagger','jira','confluence','figma','terraform','ansible',
    'jenkins','ci/cd','cicd','agile','scrum','microservices','api','oauth','jwt',
    'firebase','supabase','vercel','netlify','heroku','excel','powerbi','tableau',
    'hadoop','spark','kafka','elasticsearch','nginx','apache','c++','c#','golang',
    'rust','kotlin','swift','flutter','dart','r','matlab','scala','perl','php',
    'ruby','rails','nextjs','nuxt','svelte','tailwind','bootstrap','sass','webpack',
    'vite','redux','mobx','graphql','prisma','sequelize','typeorm','celery',
    'rabbitmq','socket.io','websocket','grpc','protobuf','openapi','langchain',
    'openai','huggingface','llm','nlp','cv','machine learning','deep learning',
    'data science','data engineering','devops','mlops','cloud','blockchain',
    'solidity','web3','linux','unix','powershell','autocad','solidworks',
    'embedded','arduino','raspberry','iot','fpga','verilog','vhdl'
}

BAD_WORDS = {
    'developers','developer','f2f','face','period','senior','lead','mandatory',
    'expertise','overview','preferred','bengaluru','bangalore','mumbai','delhi',
    'hyderabad','chennai','pune','kolkata','india','remote','hybrid','onsite',
    'immediate','joiners','joining','experience','years','year','month','months',
    'interview','hiring','actively','looking','required','must','good','strong',
    'knowledge','ability','team','work','role','position','job','apply','mode',
    'location','office','wfh','wfo','notice','important','urgent','opportunity',
    'opening','vacancy','profile','candidate','candidates','requirement',
    'requirements','skills','skill','plus','hands','working','ability','communication',
    'problem','solving','analytical','thinking','management','leadership','english',
    'verbal','written','interpersonal','client','stakeholder','cross','functional'
}

def is_valid_skill(s: str) -> bool:
    if not s or len(s) < 2 or len(s) > 40:
        return False
    s_lower = s.lower()
    if s_lower in TECH_WHITELIST:
        return True
    if any(ch in s for ch in ['.', '+', '#', '/', '-']) and len(s) < 20:
        return True
    if len(s) <= 4 and s.isupper():
        return True
    return False

def filter_skills(skills: list) -> list:
    if not skills:
        return []
    filtered, seen = [], set()
    for skill in skills:
        if not skill:
            continue
        s = str(skill).strip()
        s_lower = s.lower()
        if s_lower in seen or s_lower in BAD_WORDS:
            continue
        if is_valid_skill(s):
            filtered.append(s)
            seen.add(s_lower)
    return filtered


# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────

def clean(text) -> str:
    """Strip non-ASCII / control chars that cause PDF black squares."""
    if not text:
        return ""
    cleaned = ""
    for ch in str(text):
        cat = unicodedata.category(ch)
        if cat.startswith('C') and ch not in ('\n', '\t', ' '):
            cleaned += ' '
        elif ord(ch) > 127:
            # Replace common unicode chars with ASCII equivalents
            replacements = {
                '\u2022': '-', '\u2023': '-', '\u25cf': '-', '\u2013': '-',
                '\u2014': '-', '\u2018': "'", '\u2019': "'", '\u201c': '"',
                '\u201d': '"', '\u2026': '...', '\u00e9': 'e', '\u00e0': 'a',
            }
            cleaned += replacements.get(ch, '')
        else:
            cleaned += ch
    return cleaned.strip()

def extract_keywords(text: str) -> set:
    words = re.findall(r'\b[a-zA-Z][a-zA-Z0-9+#./-]{1,20}\b', text.lower())
    stop = {'the','and','for','with','that','this','from','your','have','will',
            'our','all','are','can','not','but','job','work','team','role'}
    return {w for w in words if w not in stop and len(w) > 2}

def parse_ai_response(raw: str) -> dict:
    """Extract JSON from AI response, handling markdown fences."""
    raw = raw.strip()
    # Remove markdown fences
    if '```' in raw:
        parts = raw.split('```')
        for part in parts:
            part = part.strip().lstrip('json').strip()
            if part.startswith('{'):
                raw = part
                break
    # Find first { to last }
    start = raw.find('{')
    end = raw.rfind('}')
    if start == -1 or end == -1:
        return {}
    try:
        return json.loads(raw[start:end+1])
    except:
        return {}

def clean_contact(contact):
    """Format contact line, removing duplicates."""
    if not contact:
        return ""
    parts = [p.strip() for p in str(contact).replace('|', ' | ').split('|')]
    seen, unique = set(), []
    for p in parts:
        key = re.sub(r'\s+', '', p.lower())
        if key and key not in seen:
            unique.append(p)
            seen.add(key)
    return ' | '.join(unique)


# ─────────────────────────────────────────────
# PDF GENERATOR
# ─────────────────────────────────────────────

def generate_tailored_pdf(tailored: dict, job_title: str, company: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
        rightMargin=18*mm, leftMargin=18*mm,
        topMargin=14*mm, bottomMargin=14*mm)
    story = []

    # ── Styles ──────────────────────────────
    name_s    = ParagraphStyle('n', fontSize=18, fontName='Helvetica-Bold',
                    textColor=colors.HexColor('#1a1a2e'), spaceAfter=3, alignment=TA_CENTER)
    contact_s = ParagraphStyle('c', fontSize=9, fontName='Helvetica',
                    textColor=colors.HexColor('#555555'), spaceAfter=4, alignment=TA_CENTER)
    link_s    = ParagraphStyle('lk', fontSize=9, fontName='Helvetica',
                    textColor=colors.HexColor('#7c6ff7'), spaceAfter=6, alignment=TA_CENTER)
    tag_s     = ParagraphStyle('t', fontSize=8.5, fontName='Helvetica',
                    textColor=colors.HexColor('#7c6ff7'), spaceAfter=8, alignment=TA_CENTER)
    sec_s     = ParagraphStyle('sec', fontSize=10.5, fontName='Helvetica-Bold',
                    textColor=colors.HexColor('#1a1a2e'), spaceBefore=8, spaceAfter=2)
    body_s    = ParagraphStyle('b', fontSize=9.5, fontName='Helvetica',
                    textColor=colors.HexColor('#333333'), spaceAfter=3, leading=14,
                    alignment=TA_JUSTIFY)
    bullet_s  = ParagraphStyle('bl', fontSize=9.5, fontName='Helvetica',
                    textColor=colors.HexColor('#333333'), leftIndent=10, spaceAfter=2,
                    leading=13, bulletIndent=0)
    job_title_s = ParagraphStyle('jt', fontSize=10, fontName='Helvetica-Bold',
                    textColor=colors.HexColor('#1a1a2e'), spaceAfter=1)
    job_sub_s = ParagraphStyle('js', fontSize=9, fontName='Helvetica',
                    textColor=colors.HexColor('#666666'), spaceAfter=3)
    proj_s    = ParagraphStyle('pr', fontSize=10, fontName='Helvetica-Bold',
                    textColor=colors.HexColor('#1a1a2e'), spaceAfter=1)

    def section(title):
        story.append(Paragraph(clean(title), sec_s))
        story.append(HRFlowable(width="100%", thickness=0.5,
                                color=colors.HexColor('#7c6ff7'), spaceAfter=4))

    # ── Header ──────────────────────────────
    name = clean(tailored.get('name', 'Candidate'))
    story.append(Paragraph(name, name_s))

    contact = clean_contact(clean(tailored.get('contact', '')))
    if contact:
        story.append(Paragraph(contact, contact_s))

    # LinkedIn & GitHub on same line
    linkedin = clean(tailored.get('linkedin', ''))
    github   = clean(tailored.get('github', ''))
    links_parts = []
    if linkedin:
        links_parts.append(f'LinkedIn: {linkedin}')
    if github:
        links_parts.append(f'GitHub: {github}')
    if links_parts:
        story.append(Paragraph(' | '.join(links_parts), link_s))

    # Job tag
    story.append(Paragraph(f"Tailored for: {clean(job_title)} @ {clean(company)}", tag_s))
    story.append(HRFlowable(width="100%", thickness=1.2,
                            color=colors.HexColor('#7c6ff7'), spaceAfter=6))

    # ── Summary ─────────────────────────────
    summary = clean(tailored.get('summary', ''))
    if summary:
        section("PROFESSIONAL SUMMARY")
        story.append(Paragraph(summary, body_s))

    # ── Skills ──────────────────────────────
    skills = filter_skills(tailored.get('skills', []))
    if skills:
        section("TECHNICAL SKILLS")
        # 3-column table
        cols = 3
        rows = [skills[i:i+cols] for i in range(0, len(skills), cols)]
        # Pad last row
        while len(rows[-1]) < cols:
            rows[-1].append('')
        skill_style = ParagraphStyle('sk', fontSize=9, fontName='Helvetica',
            textColor=colors.HexColor('#333333'))
        table_data = [[Paragraph(f'• {clean(c)}', skill_style) for c in row] for row in rows]
        col_w = (A4[0] - 36*mm) / cols
        t = Table(table_data, colWidths=[col_w]*cols)
        t.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 2),
            ('RIGHTPADDING', (0,0), (-1,-1), 2),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ]))
        story.append(t)

    # ── Experience ──────────────────────────
    experience = tailored.get('experience', [])
    if experience:
        section("WORK EXPERIENCE")
        for exp in experience:
            if isinstance(exp, dict):
                title_text = clean(exp.get('title', exp.get('role', '')))
                comp_text  = clean(exp.get('company', exp.get('organization', '')))
                period     = clean(exp.get('period', exp.get('duration', exp.get('dates', ''))))
                loc        = clean(exp.get('location', ''))
                sub_parts  = [p for p in [comp_text, period, loc] if p]

                block = []
                if title_text:
                    block.append(Paragraph(title_text, job_title_s))
                if sub_parts:
                    block.append(Paragraph(' | '.join(sub_parts), job_sub_s))

                bullets = exp.get('bullets', exp.get('responsibilities', exp.get('achievements', [])))
                if isinstance(bullets, str):
                    bullets = [bullets]
                for b in bullets:
                    b_clean = clean(str(b)).lstrip('-•* ').strip()
                    if b_clean:
                        block.append(Paragraph(f'• {b_clean}', bullet_s))

                description = clean(exp.get('description', ''))
                if description and not bullets:
                    block.append(Paragraph(f'• {description}', bullet_s))

                story.append(KeepTogether(block))
                story.append(Spacer(1, 3))
            elif isinstance(exp, str):
                story.append(Paragraph(f'• {clean(exp)}', bullet_s))

    # ── Projects ────────────────────────────
    projects = tailored.get('projects', [])
    if projects:
        section("PROJECTS")
        for proj in projects:
            if isinstance(proj, dict):
                proj_name  = clean(proj.get('name', proj.get('title', '')))
                proj_tech  = clean(proj.get('tech', proj.get('technologies', proj.get('stack', ''))))
                proj_desc  = clean(proj.get('description', proj.get('details', '')))
                proj_link  = clean(proj.get('link', proj.get('url', proj.get('github', ''))))

                block = []
                if proj_name:
                    tech_suffix = f' — {proj_tech}' if proj_tech else ''
                    block.append(Paragraph(f'{proj_name}{tech_suffix}', proj_s))
                if proj_desc:
                    block.append(Paragraph(f'• {proj_desc}', bullet_s))

                bullets = proj.get('bullets', proj.get('highlights', []))
                if isinstance(bullets, str):
                    bullets = [bullets]
                for b in bullets:
                    b_clean = clean(str(b)).lstrip('-•* ').strip()
                    if b_clean:
                        block.append(Paragraph(f'• {b_clean}', bullet_s))

                if proj_link:
                    block.append(Paragraph(f'Link: {proj_link}', job_sub_s))

                story.append(KeepTogether(block))
                story.append(Spacer(1, 3))
            elif isinstance(proj, str):
                story.append(Paragraph(f'• {clean(proj)}', bullet_s))

    # ── Education ───────────────────────────
    education = tailored.get('education', [])
    if education:
        section("EDUCATION")
        for edu in education:
            if isinstance(edu, dict):
                deg   = clean(edu.get('degree', edu.get('qualification', '')))
                inst  = clean(edu.get('institution', edu.get('university', edu.get('college', ''))))
                yr    = clean(edu.get('year', edu.get('period', edu.get('graduation', ''))))
                grade = clean(edu.get('grade', edu.get('gpa', edu.get('cgpa', edu.get('percentage', '')))))
                parts = [p for p in [inst, yr, grade] if p]
                if deg:
                    story.append(Paragraph(deg, job_title_s))
                if parts:
                    story.append(Paragraph(' | '.join(parts), job_sub_s))
            elif isinstance(edu, str):
                story.append(Paragraph(f'• {clean(edu)}', bullet_s))

    # ── Certifications ──────────────────────
    certs = tailored.get('certifications', tailored.get('certificates', []))
    if certs:
        section("CERTIFICATIONS")
        for cert in certs:
            cert_text = clean(cert if isinstance(cert, str) else cert.get('name', str(cert)))
            if cert_text:
                story.append(Paragraph(f'• {cert_text}', bullet_s))

    # ── Languages ───────────────────────────
    languages = tailored.get('languages', [])
    if languages:
        section("LANGUAGES")
        if isinstance(languages, list):
            lang_text = ' | '.join([clean(l) if isinstance(l, str) else clean(l.get('language', str(l))) for l in languages])
        else:
            lang_text = clean(str(languages))
        if lang_text:
            story.append(Paragraph(lang_text, body_s))

    # ── Declaration ─────────────────────────
    declaration = clean(tailored.get('declaration', ''))
    if declaration:
        section("DECLARATION")
        story.append(Paragraph(declaration, body_s))
        story.append(Spacer(1, 10))
        story.append(Paragraph(f'{name}', body_s))

    doc.build(story)
    return buffer.getvalue()


# ─────────────────────────────────────────────
# PYDANTIC MODELS
# ─────────────────────────────────────────────

class TailorRequest(BaseModel):
    user_id: str
    job_title: str
    job_description: str
    company: str = "Company"
    resume_text: str = ""

class TailorSavedRequest(BaseModel):
    user_id: str
    job_title: str
    job_description: str
    company: str = "Company"
    resume_text: str = ""


# ─────────────────────────────────────────────
# AI PROMPT HELPER
# ─────────────────────────────────────────────

def build_tailor_prompt(resume_text: str, job_title: str, job_description: str, missing_kws: str) -> str:
    return f"""You are an expert resume writer. Extract ALL sections from the resume and return as JSON.

RESUME TEXT:
{resume_text[:3500]}

JOB: {job_title}
JD KEYWORDS TO ADD: {missing_kws}

CRITICAL INSTRUCTIONS:
1. Extract linkedin URL - look for linkedin.com/in/ in the text or EXTRACTED_LINKS section
2. Extract github URL - look for github.com/ in the text or EXTRACTED_LINKS section
3. Extract ALL projects with name, tech stack, description
4. Keep ALL experience bullets (do not shorten)
5. summary: Write 2-3 sentences tailored for {job_title}
6. skills: existing skills + real tech keywords from JD only (never add non-tech words)
7. Return ONLY valid JSON, no markdown, no extra text

Return this exact JSON structure:
{{
  "name": "full name",
  "contact": "phone | email | city",
  "linkedin": "full linkedin URL or empty string",
  "github": "full github URL or empty string",
  "summary": "tailored professional summary for {job_title}",
  "skills": ["Skill1", "Skill2"],
  "experience": [
    {{
      "title": "Job Title",
      "company": "Company Name",
      "period": "Month Year - Month Year",
      "location": "City",
      "bullets": ["bullet 1", "bullet 2"]
    }}
  ],
  "projects": [
    {{
      "name": "Project Name",
      "tech": "React, Python, etc",
      "description": "What it does",
      "link": "github/project URL or empty string"
    }}
  ],
  "education": [
    {{
      "degree": "Degree Name",
      "institution": "University Name",
      "year": "Year",
      "grade": "Grade/GPA"
    }}
  ],
  "certifications": ["Cert 1", "Cert 2"],
  "languages": ["English", "etc"],
  "declaration": "I hereby declare that the above information is true and correct.",
  "keywords_added": ["kw1", "kw2"]
}}"""


# ─────────────────────────────────────────────
# ENDPOINTS
# ─────────────────────────────────────────────

@router.post("/tailor")
async def tailor_resume(
    resume_file: UploadFile = File(...),
    job_title: str = Form(...),
    job_description: str = Form(...),
    company: str = Form(default="Company"),
    user_id: str = Form(...)
):
    try:
        resume_bytes = await resume_file.read()
        resume_text = extract_pdf_text(resume_bytes)

        if not resume_text or len(resume_text) < 30:
            raise HTTPException(status_code=400, detail="Could not read PDF content.")

        resume_kws = extract_keywords(resume_text)
        jd_kws = extract_keywords(job_description + " " + job_title)
        missing = jd_kws - resume_kws
        original_score = min(95, int((len(resume_kws & jd_kws) / max(len(jd_kws), 1)) * 100))
        missing_kws_str = ', '.join(list(missing)[:15])

        prompt = build_tailor_prompt(resume_text, job_title, job_description, missing_kws_str)

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[
                {"role": "system", "content": "You are a JSON API. Return only valid JSON. Never add non-technical words to skills."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,
            max_tokens=4000
        )

        raw = response.choices[0].message.content.strip()
        print(f"RAW AI (tailor): {raw[:300]}")
        tailored = parse_ai_response(raw)

        if not tailored:
            raise HTTPException(status_code=500, detail="AI response parsing failed.")

        tailored['skills'] = filter_skills(tailored.get('skills', []))
        keywords_added = filter_skills(tailored.get('keywords_added', []))

        tailored_kws = extract_keywords(json.dumps(tailored))
        new_matched = tailored_kws & jd_kws
        tailored_score = min(97, int((len(new_matched) / max(len(jd_kws), 1)) * 100))
        if tailored_score <= original_score:
            tailored_score = min(97, original_score + len(keywords_added) * 2)

        pdf_bytes = generate_tailored_pdf(tailored, job_title, company)
        filename = f"tailored_{job_title.replace(' ', '_')}.pdf"

        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "x-keywords-added": json.dumps(keywords_added),
                "x-original-score": str(original_score),
                "x-tailored-score": str(tailored_score),
                "Access-Control-Expose-Headers": "x-keywords-added,x-original-score,x-tailored-score"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"TAILOR ERROR: {traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/save-resume")
async def save_resume(
    resume_file: UploadFile = File(...),
    user_id: str = Form(...)
):
    try:
        file_bytes = await resume_file.read()
        storage_path = f"{user_id}/resume.pdf"
        supabase.storage.from_("resumes").upload(
            storage_path, file_bytes,
            {"content-type": "application/pdf", "upsert": "true"}
        )
        supabase.table("user_resumes").upsert({
            "user_id": user_id,
            "file_name": resume_file.filename,
            "file_size": len(file_bytes),
            "storage_path": storage_path,
        }, on_conflict="user_id").execute()
        return {"success": True, "file_name": resume_file.filename, "file_size": len(file_bytes)}
    except Exception as e:
        return {"success": False, "error": str(e)}


@router.get("/resume-info/{user_id}")
async def get_resume_info(user_id: str):
    try:
        result = supabase.table("user_resumes").select("*").eq("user_id", user_id).execute()
        return result.data[0] if result.data else None
    except:
        return None


@router.get("/resume-text/{user_id}")
async def get_resume_text(user_id: str):
    try:
        storage_path = f"{user_id}/resume.pdf"
        file_bytes = supabase.storage.from_("resumes").download(storage_path)
        text = extract_pdf_text(file_bytes)
        return {"text": text, "success": True}
    except Exception as e:
        return {"text": "", "success": False, "error": str(e)}


@router.post("/tailor-saved")
async def tailor_saved_resume(data: TailorSavedRequest):
    try:
        resume_text = data.resume_text

        if not resume_text:
            try:
                storage_path = f"{data.user_id}/resume.pdf"
                file_bytes = supabase.storage.from_("resumes").download(storage_path)
                resume_text = extract_pdf_text(file_bytes)
            except Exception as e:
                print(f"Storage download error: {e}")

        if not resume_text:
            raise HTTPException(status_code=400, detail="No resume found. Please upload your resume first.")

        resume_kws = extract_keywords(resume_text)
        jd_kws = extract_keywords(data.job_description + " " + data.job_title)
        missing = jd_kws - resume_kws
        original_score = min(95, int((len(resume_kws & jd_kws) / max(len(jd_kws), 1)) * 100))
        missing_kws_str = ', '.join(list(missing)[:15])

        prompt = build_tailor_prompt(resume_text, data.job_title, data.job_description, missing_kws_str)

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[
                {"role": "system", "content": "You are a JSON API. Return only valid JSON. Never add non-technical words to skills."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,
            max_tokens=4000
        )

        raw = response.choices[0].message.content.strip()
        print(f"RAW AI (tailor-saved): {raw[:300]}")
        tailored = parse_ai_response(raw)

        if not tailored:
            raise HTTPException(status_code=500, detail="AI response parsing failed.")

        tailored['skills'] = filter_skills(tailored.get('skills', []))
        keywords_added = filter_skills(tailored.get('keywords_added', []))

        tailored_kws = extract_keywords(json.dumps(tailored))
        new_matched = tailored_kws & jd_kws
        tailored_score = min(97, int((len(new_matched) / max(len(jd_kws), 1)) * 100))
        if tailored_score <= original_score:
            tailored_score = min(97, original_score + len(keywords_added) * 2)

        pdf_bytes = generate_tailored_pdf(tailored, data.job_title, data.company)
        filename = f"tailored_{data.job_title.replace(' ', '_')}.pdf"

        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "x-keywords-added": json.dumps(keywords_added),
                "x-original-score": str(original_score),
                "x-tailored-score": str(tailored_score),
                "Access-Control-Expose-Headers": "x-keywords-added,x-original-score,x-tailored-score"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"TAILOR-SAVED ERROR: {traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/tailor-docx")
async def tailor_resume_docx(
    resume_file: UploadFile = File(...),
    job_title: str = Form(...),
    job_description: str = Form(...),
    company: str = Form(default="Company"),
    user_id: str = Form(...)
):
    """DOCX generation — returns same tailored data as DOCX via python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        resume_bytes = await resume_file.read()
        resume_text = extract_pdf_text(resume_bytes)

        if not resume_text or len(resume_text) < 30:
            raise HTTPException(status_code=400, detail="Could not read PDF.")

        resume_kws = extract_keywords(resume_text)
        jd_kws = extract_keywords(job_description + " " + job_title)
        missing = jd_kws - resume_kws
        original_score = min(95, int((len(resume_kws & jd_kws) / max(len(jd_kws), 1)) * 100))
        missing_kws_str = ', '.join(list(missing)[:15])

        prompt = build_tailor_prompt(resume_text, job_title, job_description, missing_kws_str)

        response = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[
                {"role": "system", "content": "You are a JSON API. Return only valid JSON."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,
            max_tokens=4000
        )

        raw = response.choices[0].message.content.strip()
        tailored = parse_ai_response(raw)
        tailored['skills'] = filter_skills(tailored.get('skills', []))
        keywords_added = filter_skills(tailored.get('keywords_added', []))

        # Build DOCX
        doc = Document()
        # Margins
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        def add_heading(text, size=14, bold=True, center=True, color=None):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean(text))
            run.bold = bold
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor(*color)
            return p

        def add_section(title):
            p = doc.add_paragraph()
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(26, 26, 46)
            p.paragraph_format.space_before = Pt(8)
            # Underline via border would need XML manipulation, skip for now

        # Name
        add_heading(tailored.get('name', 'Candidate'), size=18, color=(26, 26, 46))
        # Contact
        contact = clean_contact(clean(tailored.get('contact', '')))
        if contact:
            p = doc.add_paragraph(contact)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Links
        linkedin = clean(tailored.get('linkedin', ''))
        github = clean(tailored.get('github', ''))
        links = ' | '.join(filter(None, [
            f'LinkedIn: {linkedin}' if linkedin else '',
            f'GitHub: {github}' if github else ''
        ]))
        if links:
            p = doc.add_paragraph(links)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Summary
        summary = clean(tailored.get('summary', ''))
        if summary:
            add_section('PROFESSIONAL SUMMARY')
            doc.add_paragraph(summary)

        # Skills
        skills = tailored.get('skills', [])
        if skills:
            add_section('TECHNICAL SKILLS')
            doc.add_paragraph(' | '.join(skills))

        # Experience
        experience = tailored.get('experience', [])
        if experience:
            add_section('WORK EXPERIENCE')
            for exp in experience:
                if isinstance(exp, dict):
                    p = doc.add_paragraph()
                    run = p.add_run(clean(exp.get('title', exp.get('role', ''))))
                    run.bold = True
                    doc.add_paragraph(clean(f"{exp.get('company','')} | {exp.get('period','')} | {exp.get('location','')}"))
                    for b in exp.get('bullets', exp.get('responsibilities', [])):
                        doc.add_paragraph(f'• {clean(str(b))}')

        # Projects
        projects = tailored.get('projects', [])
        if projects:
            add_section('PROJECTS')
            for proj in projects:
                if isinstance(proj, dict):
                    name_tech = f"{proj.get('name','')} — {proj.get('tech','')}" if proj.get('tech') else proj.get('name','')
                    p = doc.add_paragraph()
                    run = p.add_run(clean(name_tech))
                    run.bold = True
                    if proj.get('description'):
                        doc.add_paragraph(f'• {clean(proj["description"])}')
                    for b in proj.get('bullets', []):
                        doc.add_paragraph(f'• {clean(str(b))}')
                    if proj.get('link'):
                        doc.add_paragraph(f'Link: {clean(proj["link"])}')

        # Education
        education = tailored.get('education', [])
        if education:
            add_section('EDUCATION')
            for edu in education:
                if isinstance(edu, dict):
                    p = doc.add_paragraph()
                    run = p.add_run(clean(edu.get('degree', '')))
                    run.bold = True
                    doc.add_paragraph(clean(f"{edu.get('institution','')} | {edu.get('year','')} | {edu.get('grade','')}"))

        # Certifications
        certs = tailored.get('certifications', [])
        if certs:
            add_section('CERTIFICATIONS')
            for c in certs:
                doc.add_paragraph(f'• {clean(c if isinstance(c, str) else c.get("name", str(c)))}')

        # Languages
        langs = tailored.get('languages', [])
        if langs:
            add_section('LANGUAGES')
            doc.add_paragraph(' | '.join([clean(l) if isinstance(l, str) else clean(l.get('language', str(l))) for l in langs]))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        filename = f"tailored_{job_title.replace(' ', '_')}.docx"

        tailored_score = min(97, original_score + len(keywords_added) * 2)

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "x-keywords-added": json.dumps(keywords_added),
                "x-original-score": str(original_score),
                "x-tailored-score": str(tailored_score),
                "Access-Control-Expose-Headers": "x-keywords-added,x-original-score,x-tailored-score"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"TAILOR-DOCX ERROR: {traceback.format_exc()}")
        raise HTTPException(status_code=500, detail=str(e))